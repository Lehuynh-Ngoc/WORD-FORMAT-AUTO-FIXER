import os
import json
import logging
import urllib.request
import urllib.error
from typing import Optional, List, Dict
from dataclasses import dataclass, field

from src.core.ai_config import get_active_provider, ProviderConfig, APIStatus

logger = logging.getLogger("wordfix")


@dataclass
class Question:
    id: int
    text: str
    hint: str = ""
    answer: str = ""


@dataclass
class GenerateResult:
    title: str = ""
    content_markdown: str = ""
    output_path: str = ""
    questions_asked: int = 0


def _call_ai(provider: ProviderConfig, system_prompt: str, user_prompt: str) -> str:
    if provider.name in ("openai", "deepseek"):
        return _call_openai_compat(provider, system_prompt, user_prompt)
    elif provider.name == "gemini":
        return _call_gemini(provider, system_prompt, user_prompt)
    elif provider.name == "claude":
        return _call_claude(provider, system_prompt, user_prompt)
    else:
        # Moi custom provider (ten tuy y) deu dung OpenAI-compatible API
        return _call_openai_compat(provider, system_prompt, user_prompt)


def _call_openai_compat(provider: ProviderConfig, system_prompt: str, user_prompt: str) -> str:
    payload = json.dumps({
        "model": provider.model,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.7,
    }).encode()

    req = urllib.request.Request(
        f"{provider.base_url}/chat/completions",
        data=payload,
        headers={
            "Authorization": f"Bearer {provider.api_key}",
            "Content-Type": "application/json",
        },
    )
    resp = urllib.request.urlopen(req, timeout=120)
    data = json.loads(resp.read().decode())
    return data["choices"][0]["message"]["content"].strip()


def _call_gemini(provider: ProviderConfig, system_prompt: str, user_prompt: str) -> str:
    full_prompt = f"{system_prompt}\n\n{user_prompt}"
    payload = json.dumps({
        "contents": [{"parts": [{"text": full_prompt}]}],
        "generationConfig": {"temperature": 0.7},
    }).encode()

    req = urllib.request.Request(
        f"{provider.base_url}/models/{provider.model}:generateContent?key={provider.api_key}",
        data=payload,
        headers={"Content-Type": "application/json"},
    )
    resp = urllib.request.urlopen(req, timeout=120)
    data = json.loads(resp.read().decode())
    return data["candidates"][0]["content"]["parts"][0]["text"].strip()


def _call_claude(provider: ProviderConfig, system_prompt: str, user_prompt: str) -> str:
    payload = json.dumps({
        "model": provider.model,
        "max_tokens": 4096,
        "system": system_prompt,
        "messages": [{"role": "user", "content": user_prompt}],
        "temperature": 0.7,
    }).encode()

    req = urllib.request.Request(
        f"{provider.base_url}/messages",
        data=payload,
        headers={
            "x-api-key": provider.api_key,
            "anthropic-version": "2023-06-01",
            "Content-Type": "application/json",
        },
    )
    resp = urllib.request.urlopen(req, timeout=120)
    data = json.loads(resp.read().decode())
    return data["content"][0]["text"].strip()


def ask_ai_questions(doc_type: str, provider: ProviderConfig) -> List[Question]:
    system_prompt = (
        "You are a professional Vietnamese document assistant. "
        "When the user tells you what kind of document they need, "
        "you must ask them questions to gather all the necessary information "
        "to write that document.\n\n"
        "RULES:\n"
        "- Ask clear, specific questions in Vietnamese\n"
        "- Ask 4-8 questions depending on the document type\n"
        "- Include a 'hint' for each question explaining what information is needed\n"
        "- Return ONLY a JSON array, no other text\n"
        "- Each item: {\"id\": 1, \"text\": \"...\", \"hint\": \"...\"}\n"
        "- Make questions natural and friendly\n"
        "- Cover: names, dates, reasons, addresses, phone numbers etc. as appropriate"
    )

    user_prompt = (
        f"I need to create a document of type: \"{doc_type}\"\n"
        f"Please ask me the necessary questions to gather information for this document."
    )

    response = _call_ai(provider, system_prompt, user_prompt)

    try:
        json_str = response
        if "```json" in json_str:
            json_str = json_str.split("```json")[1].split("```")[0]
        elif "```" in json_str:
            json_str = json_str.split("```")[1].split("```")[0]
        json_str = json_str.strip()

        if not json_str.startswith("["):
            start = json_str.find("[")
            end = json_str.rfind("]") + 1
            if start >= 0 and end > start:
                json_str = json_str[start:end]

        questions_data = json.loads(json_str)
        return [
            Question(id=q["id"], text=q["text"], hint=q.get("hint", ""))
            for q in questions_data
        ]
    except (json.JSONDecodeError, KeyError) as e:
        logger.error(f"Failed to parse AI response as questions: {e}")
        logger.debug(f"Raw response: {response[:500]}")
        return []


def generate_document_content(
    doc_type: str,
    questions: List[Question],
    provider: ProviderConfig,
) -> tuple:
    qa_text = "\n".join(
        f"Câu hỏi: {q.text}\nTrả lời: {q.answer}"
        for q in questions if q.answer
    )

    system_prompt = (
        "You are a professional Vietnamese document writer. "
        "Create a well-formatted document based on the information provided.\n\n"
        "RULES:\n"
        "- Write in formal Vietnamese\n"
        "- Use proper document structure with headings and paragraphs\n"
        "- Include all information from the answers\n"
        "- Return the document in this EXACT markdown format:\n"
        "  # TITLE (document title)\n"
        "  ## Heading sections as needed\n"
        "  Normal paragraphs for body text\n"
        "- The document must be complete and ready to use\n"
        "- Do NOT include placeholder text like [...]\n"
        "- Format: use # for title, ## for section headers, plain text for paragraphs\n"
        "- Each paragraph should be on its own line separated by blank lines"
    )

    user_prompt = (
        f"Create a \"{doc_type}\" document with the following information:\n\n"
        f"{qa_text}\n\n"
        f"Return the complete document content in markdown format."
    )

    response = _call_ai(provider, system_prompt, user_prompt)

    title = doc_type
    lines = response.strip().split("\n")
    if lines and lines[0].startswith("#"):
        title = lines[0].lstrip("#").strip()

    return title, response.strip()


def save_as_docx(title: str, content_markdown: str, output_path: str) -> str:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    import re as _re
    heading_re = _re.compile(r'^(#{1,4})\s+(.*)$')

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(13)

    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)

    lines = content_markdown.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        heading_match = heading_re.match(line)

        if heading_match and len(heading_match.group(1)) == 1:
            heading_text = heading_match.group(2).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(heading_text)
            run.bold = True
            run.font.size = Pt(16)
            run.font.name = "Times New Roman"
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)

        elif heading_match and len(heading_match.group(1)) == 2:
            heading_text = heading_match.group(2).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(heading_text)
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = "Times New Roman"
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)

        elif heading_match:
            heading_text = heading_match.group(2).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(heading_text)
            run.bold = True
            run.font.size = Pt(13)
            run.font.name = "Times New Roman"
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)

        elif line.startswith("- "):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(bullet_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)

        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)

        i += 1

    output_dir = os.path.dirname(os.path.abspath(output_path))
    if output_dir and not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)

    doc.save(output_path)
    return output_path

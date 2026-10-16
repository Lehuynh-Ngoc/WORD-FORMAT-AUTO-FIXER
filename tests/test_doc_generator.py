import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document

from src.core.doc_generator import save_as_docx


def load_paragraphs(path):
    doc = Document(path)
    return [(p.text, p.runs) for p in doc.paragraphs]


def test_save_as_docx_heading_levels(tmp_path):
    """Bug #12: #, ##, ###, #### deu phai tao heading, khong de thuong."""
    content = (
        "# Tieu de chinh\n\n"
        "## Muc 1\n\n"
        "### Muc 1.1\n\n"
        "#### Muc 1.1.1\n\n"
        "Doan van binh thuong.\n\n"
        "- item mot\n"
        "- item hai\n"
    )
    out = str(tmp_path / "out.docx")
    save_as_docx("Test", content, out)

    paras = load_paragraphs(out)
    texts = [t for t, _ in paras]

    assert "Tieu de chinh" in texts
    assert "Muc 1" in texts
    assert "Muc 1.1" in texts
    assert "Muc 1.1.1" in texts
    assert "#### Muc 1.1.1" not in texts  # khong con chu # thua

    # Heading 1 (dau tien) phai co size 16, bold
    h1_runs = paras[0][1]
    assert h1_runs[0].bold is True
    assert h1_runs[0].font.size.pt == 16

    # Heading 2 size 14
    h2_runs = paras[1][1]
    assert h2_runs[0].font.size.pt == 14

    # Heading 3 & 4 size 13
    assert paras[2][1][0].font.size.pt == 13
    assert paras[3][1][0].font.size.pt == 13


def test_save_as_docx_preserves_inner_spaces(tmp_path):
    """lstrip cu nuot ca space o dau heading."""
    content = "#    Tieu de co space dau"
    out = str(tmp_path / "out.docx")
    save_as_docx("Test", content, out)

    texts = [t for t, _ in load_paragraphs(out)]
    assert "Tieu de co space dau" in texts


def test_save_as_docx_bullet(tmp_path):
    content = "- noi dung gach dau"
    out = str(tmp_path / "out.docx")
    save_as_docx("Test", content, out)

    doc = Document(out)
    assert doc.paragraphs[0].style.name == "List Bullet"
    assert doc.paragraphs[0].text == "noi dung gach dau"


def test_save_as_docx_creates_output_dir(tmp_path):
    out_dir = tmp_path / "sub" / "dir"
    out = str(out_dir / "out.docx")

    result = save_as_docx("Tieu de", "Noi dung", out)

    assert result == out
    assert os.path.exists(out)


def test_save_as_docx_normal_paragraph_justified(tmp_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    out = str(tmp_path / "out.docx")
    save_as_docx("Tieu de", "Doan van thuong", out)

    doc = Document(out)
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert doc.paragraphs[0].runs[0].font.size.pt == 13

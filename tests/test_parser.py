import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document
from docx.shared import Pt

from src.parsers.docx_parser import DocxParser, iter_document_paragraphs


def make_docx_with_table(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("doan van tren bang")
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "o mot"
    table.cell(0, 1).text = "o hai"
    table.cell(1, 0).text = "o ba"
    table.cell(1, 1).text = "o bon"

    doc.add_paragraph("doan van duoi bang")

    path = str(tmp_path / "table.docx")
    doc.save(path)
    return path


def test_read_includes_table_paragraphs(tmp_path):
    path = make_docx_with_table(tmp_path)
    parsed = DocxParser().read(path)

    texts = [p.text for p in parsed.paragraphs]
    assert "doan van tren bang" in texts
    assert "o mot" in texts
    assert "o hai" in texts
    assert "o ba" in texts
    assert "o bon" in texts
    assert "doan van duoi bang" in texts


def test_read_order_body_table_body(tmp_path):
    path = make_docx_with_table(tmp_path)
    parsed = DocxParser().read(path)

    texts = [p.text for p in parsed.paragraphs]
    i_body1 = texts.index("doan van tren bang")
    i_cell = texts.index("o mot")
    i_body2 = texts.index("doan van duoi bang")
    assert i_body1 < i_cell < i_body2


def test_table_paragraph_source_reference(tmp_path):
    """Moi ParsedParagraph trong table phai tro dung vao paragraph goc."""
    path = make_docx_with_table(tmp_path)
    parsed = DocxParser().read(path)

    for pp in parsed.paragraphs:
        assert pp.source is not None


def test_save_writes_back_to_table_cell(tmp_path):
    """Sua text trong ParsedParagraph cua table cell phai ghi vao dung cell."""
    path = make_docx_with_table(tmp_path)
    parser = DocxParser()
    parsed = parser.read(path)

    target = next(p for p in parsed.paragraphs if p.text == "o ba")
    target.runs[0].text = "da sua"
    if not target.runs:
        target.source.text = "da sua"

    out = str(tmp_path / "out.docx")
    parser.save(parsed, out)

    saved = Document(out)
    all_texts = [p.text for p in iter_document_paragraphs(saved)]
    assert "da sua" in all_texts
    assert "o ba" not in all_texts


def test_save_roundtrip_preserves_other_cells(tmp_path):
    path = make_docx_with_table(tmp_path)
    parser = DocxParser()
    parsed = parser.read(path)

    out = str(tmp_path / "out.docx")
    parser.save(parsed, out)

    saved = Document(out)
    texts = [p.text for p in iter_document_paragraphs(saved)]
    for expected in ["o mot", "o hai", "o bon", "doan van tren bang", "doan van duoi bang"]:
        assert expected in texts


def test_merged_cell_counted_once():
    """Cell merge xuat hien nhieu lan trong row.cells — phai dedup."""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "merge"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"

    paras = list(iter_document_paragraphs(doc))
    texts = [p.text for p in paras]
    assert texts.count("merge") == 1


def test_count_includes_table_words(tmp_path):
    from src.core.counter import count_document

    path = make_docx_with_table(tmp_path)
    result = count_document(path)

    # 2 doan body (4 tu + 4 tu) + 4 cell (2 tu moi cell: "o mot"...) = 16 tu
    assert result.words == 16


def test_save_writes_changed_text_back(tmp_path):
    """Bug translate: save() truoc day chi ghi style/font, text sua/bi dich khong
    duoc ghi vao file — file ra van giong ban dau."""
    doc = Document()
    p = doc.add_paragraph()
    run1 = p.add_run("xin ")
    run2 = p.add_run("chao")
    path = str(tmp_path / "in.docx")
    doc.save(path)

    parser = DocxParser()
    parsed = parser.read(path)

    # Mo phong dich: text chuyen sang doan moi, cac run cuoc bi xoa
    parsed.paragraphs[0].runs[0].text = "hello"
    parsed.paragraphs[0].runs[1].text = ""

    out = str(tmp_path / "out.docx")
    parser.save(parsed, out)

    saved = Document(out)
    assert saved.paragraphs[0].text == "hello"


def test_save_fallback_without_source_reference(tmp_path):
    """ParsedDocument tao thu cong (khong co source) van save duoc theo thu tu."""
    from src.parsers.base import ParsedDocument, ParsedParagraph, ParsedRun

    path = make_docx_with_table(tmp_path)
    parser = DocxParser()
    parsed = parser.read(path)

    # Bo tham chiếu source — ép fallback theo index
    for pp in parsed.paragraphs:
        pp.source = None

    out = str(tmp_path / "out.docx")
    parser.save(parsed, out)

    saved = Document(out)
    texts = [p.text for p in iter_document_paragraphs(saved)]
    assert "o mot" in texts
    assert "doan van duoi bang" in texts

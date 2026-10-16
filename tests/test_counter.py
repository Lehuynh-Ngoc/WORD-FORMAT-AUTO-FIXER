import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document

from src.core.counter import count_document, _format_duration, _detect_languages


def make_docx(tmp_path, paragraphs):
    """paragraphs: list of (style_name, [run_texts])."""
    doc = Document()
    # Xoa doan trong mac dinh cua Document moi
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    for style_name, runs in paragraphs:
        p = doc.add_paragraph(style=style_name)
        for text in runs:
            p.add_run(text)

    path = str(tmp_path / "test.docx")
    doc.save(path)
    return path


def test_heading_counted_once_per_paragraph(tmp_path):
    """Bug #4: heading co nhieu run phai van chi dem 1 lan."""
    path = make_docx(tmp_path, [
        ("Heading 1", ["Chuong 1", " - ", "Mo dau"]),   # 3 runs
        ("Heading 1", ["Ket luan"]),                     # 1 run
        ("Heading 2", ["1.1", " Noi dung"]),             # 2 runs
        ("Normal", ["Doan van binh thuong."]),
    ])

    result = count_document(path)

    assert result.headings["Heading 1"] == 2
    assert result.headings["Heading 2"] == 1
    assert result.paragraphs == 4


def test_word_and_char_counts(tmp_path):
    path = make_docx(tmp_path, [
        ("Normal", ["xin chao the gioi"]),
        ("Normal", ["mot hai ba"]),
    ])

    result = count_document(path)

    assert result.words == 7  # "xin chao the gioi" (4) + "mot hai ba" (3)
    assert result.lines == 2
    assert result.characters_with_space > 0
    assert result.pages_estimated >= 1


def test_empty_document_min_one_page(tmp_path):
    path = make_docx(tmp_path, [("Normal", [""])])

    result = count_document(path)

    assert result.pages_estimated >= 1


def test_format_duration():
    assert _format_duration(0.5) == "< 1 phut"
    assert _format_duration(3) == "3 phut"
    assert _format_duration(60) == "1 gio"
    assert _format_duration(90) == "1 gio 30 phut"


def test_detect_languages():
    assert "Tieng Viet" in _detect_languages("Đây là tiếng Việt có dấu")
    assert _detect_languages("only short english") == []
